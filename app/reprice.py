from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook.workbook import Workbook

from app.parsers import norm, parse_number, parse_discount_percent, ceil_int


# =========================
# Канонические имена колонок
# =========================

@dataclass(frozen=True)
class CanonicalCols:
    p_before: str
    discount: str
    p_current: str
    p_target: str
    price_min: str
    p_new: str


DEFAULT_CANONICAL = CanonicalCols(
    p_before="Текущая цена",
    discount="Текущая скидка",
    p_current="текущая_конечная_цена",
    p_target="желаемая_конечная_цена",
    price_min="цена_мин",
    p_new="Новая цена",
)

SYNONYMS: Dict[str, List[str]] = {
    "p_before": [
        "текущая цена", "цена", "цена продавца", "цена до скидки", "цена без скидки",
        "price", "regular price",
    ],
    "discount": [
        "текущая скидка", "скидка", "скидка продавца", "скидка %", "скидка, %",
        "seller discount",
    ],
    "p_current": [
        "текущая_конечная_цена", "текущая конечная цена", "конечная цена",
        "цена для покупателя", "цена к оплате", "к оплате", "финальная цена",
        "buyer price", "final price",
    ],
    "p_target": [
        "желаемая_конечная_цена", "желаемая конечная цена", "целевая конечная цена",
        "желаемая цена", "целевая цена", "target price",
    ],
    "price_min": [
        "цена_мин", "цена мин", "минимальная цена", "мин цена", "минимум",
        "price_min", "min price", "minimum price",
    ],
    "p_new": [
        "новая цена", "новая цена до скидки", "новая цена продавца", "price new",
    ],
}


# =========================
# Формула
# =========================

def compute_p_new(p_before: float, d_percent: float, p_current: float, p_target: float) -> Optional[int]:
    factor = 1.0 - (d_percent / 100.0)
    if factor <= 0:
        return None

    p_after = p_before * factor
    if p_after == 0:
        return None

    k = p_current / p_after
    if k == 0:
        return None

    p_new = p_target / (k * factor)
    return ceil_int(p_new)


# =========================
# Поиск заголовков
# =========================

def build_header_map(ws: Worksheet, header_row: int) -> Dict[str, int]:
    m: Dict[str, int] = {}
    for cell in ws[header_row]:
        key = norm(cell.value)
        if key:
            m[key] = cell.col_idx
    return m


def choose_col(header_map: Dict[str, int], candidates: Iterable[str]) -> Optional[int]:
    for name in candidates:
        key = norm(name)
        if key in header_map:
            return header_map[key]
    return None


def header_row_score(row_values_norm: List[str]) -> int:
    present = set(row_values_norm)
    score = 0
    for group, names in SYNONYMS.items():
        if group == "p_new":
            continue
        if any(norm(n) in present for n in names):
            score += 1
    return score


def find_best_header_row(ws: Worksheet, scan_rows: int = 50) -> Optional[int]:
    best_row = None
    best_score = -1
    max_r = min(scan_rows, ws.max_row)
    for r in range(1, max_r + 1):
        values = [norm(c.value) for c in ws[r]]
        if not any(values):
            continue
        score = header_row_score(values)
        if score > best_score:
            best_score = score
            best_row = r
    if best_score >= 4:
        return best_row
    return None


def ensure_output_column(ws: Worksheet, header_row: int, header_map: Dict[str, int]) -> int:
    existing = choose_col(header_map, SYNONYMS["p_new"])
    if existing is not None:
        return existing

    new_col = ws.max_column + 1
    ws.cell(row=header_row, column=new_col, value=DEFAULT_CANONICAL.p_new)
    header_map[norm(DEFAULT_CANONICAL.p_new)] = new_col
    return new_col


def detect_required_columns(ws: Worksheet, header_row: int) -> Tuple[int, int, int, int, int, int]:
    header_map = build_header_map(ws, header_row)

    p_before_col  = choose_col(header_map, SYNONYMS["p_before"])
    discount_col  = choose_col(header_map, SYNONYMS["discount"])
    p_current_col = choose_col(header_map, SYNONYMS["p_current"])
    p_target_col  = choose_col(header_map, SYNONYMS["p_target"])
    price_min_col = choose_col(header_map, SYNONYMS["price_min"])
    p_new_col     = ensure_output_column(ws, header_row, header_map)

    missing = []
    if p_before_col is None:  missing.append("P_before")
    if discount_col is None:  missing.append("d")
    if p_current_col is None: missing.append("P_current")
    if p_target_col is None:  missing.append("P_target")
    if price_min_col is None: missing.append("цена_мин")
    if missing:
        raise RuntimeError("Не нашёл обязательные колонки: " + ", ".join(missing))

    return p_before_col, discount_col, p_current_col, p_target_col, price_min_col, p_new_col


# =========================
# Обработка файлов
# =========================

def process_xlsx_reprice(in_path: str, out_path: str, sheet_name: str = None) -> int:
    wb: Workbook = load_workbook(in_path)
    ws: Worksheet = wb[sheet_name] if sheet_name else wb.active

    header_row = find_best_header_row(ws, scan_rows=50)
    if header_row is None:
        raise RuntimeError("Не удалось определить строку заголовков (в первых 50 строках).")

    p_before_col, discount_col, p_current_col, p_target_col, price_min_col, p_new_col = \
        detect_required_columns(ws, header_row)

    filled = 0
    for r in range(header_row + 1, ws.max_row + 1):
        p_before  = parse_number(ws.cell(r, p_before_col).value)
        d         = parse_discount_percent(ws.cell(r, discount_col).value)
        p_current = parse_number(ws.cell(r, p_current_col).value)
        p_target  = parse_number(ws.cell(r, p_target_col).value)
        p_min     = parse_number(ws.cell(r, price_min_col).value)

        if any(v is None for v in (p_before, d, p_current, p_target, p_min)):
            continue

        p_new_calc = compute_p_new(p_before, d, p_current, p_target)
        if p_new_calc is None:
            continue

        p_min_int   = ceil_int(p_min)
        p_new_final = p_new_calc if p_new_calc >= p_min_int else p_min_int

        ws.cell(row=r, column=p_new_col, value=int(p_new_final))
        filled += 1

    wb.save(out_path)
    return filled


def process_xls_to_xlsx_reprice(in_path: str, out_path: str, sheet_name: str = None) -> int:
    import pandas as pd

    df = pd.read_excel(in_path, sheet_name=sheet_name, engine="xlrd")
    if isinstance(df, dict):
        first_sheet = next(iter(df.keys()))
        df = df[first_sheet]

    col_norm = {c: norm(c) for c in df.columns}

    def find_df_col(group: str) -> Optional[str]:
        for syn in SYNONYMS[group]:
            syn_n = norm(syn)
            for c, cn in col_norm.items():
                if cn == syn_n:
                    return c
        return None

    c_p_before  = find_df_col("p_before")
    c_discount  = find_df_col("discount")
    c_p_current = find_df_col("p_current")
    c_p_target  = find_df_col("p_target")
    c_price_min = find_df_col("price_min")

    missing = []
    if c_p_before is None:  missing.append("P_before")
    if c_discount is None:  missing.append("d")
    if c_p_current is None: missing.append("P_current")
    if c_p_target is None:  missing.append("P_target")
    if c_price_min is None: missing.append("цена_мин")
    if missing:
        raise RuntimeError(f"В .xls не нашёл колонки: {', '.join(missing)}")

    out_col = None
    for syn in SYNONYMS["p_new"]:
        for c in df.columns:
            if norm(c) == norm(syn):
                out_col = c
                break
        if out_col:
            break

    if out_col is None:
        out_col = DEFAULT_CANONICAL.p_new
        df[out_col] = None

    filled = 0
    for idx, row in df.iterrows():
        p_before  = parse_number(row.get(c_p_before))
        d         = parse_discount_percent(row.get(c_discount))
        p_current = parse_number(row.get(c_p_current))
        p_target  = parse_number(row.get(c_p_target))
        p_min     = parse_number(row.get(c_price_min))

        if any(v is None for v in (p_before, d, p_current, p_target, p_min)):
            continue

        p_new_calc = compute_p_new(p_before, d, p_current, p_target)
        if p_new_calc is None:
            continue

        p_min_int   = ceil_int(p_min)
        p_new_final = p_new_calc if p_new_calc >= p_min_int else p_min_int

        df.at[idx, out_col] = int(p_new_final)
        filled += 1

    df.to_excel(out_path, index=False, engine="openpyxl")
    return filled


def reprice_process_file(input_path: str, output_path: str) -> int:
    ext = Path(input_path).suffix.lower()
    if ext == ".xlsx":
        return process_xlsx_reprice(input_path, output_path)
    if ext == ".xls":
        return process_xls_to_xlsx_reprice(input_path, output_path)
    raise RuntimeError("Поддерживаются только .xlsx и .xls")
